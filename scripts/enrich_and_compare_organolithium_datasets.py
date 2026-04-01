#!/usr/bin/env python3
"""Safely enrich the local organolithium sheet and compare it with ORD/USPTO-LLM."""

from __future__ import annotations

import csv
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib.pyplot as plt
import openpyxl
import pandas as pd
import seaborn as sns
from rdkit import Chem
from rdkit import RDLogger


ROOT = Path("/Users/zhaowenyuan/Projects/FlowFigTabMiner")
USER_XLSX = Path(
    "/Users/zhaowenyuan/Library/CloudStorage/OneDrive-共享的库-Onedrive/260329/FlowFigTabMiner/FlowFigTabMiner/data/final_output/ALL_organolithium_normalized_20260327_1022.xlsx"
)
ORD_RICH_CSV = ROOT / "other dataset" / "ord" / "organolithium" / "ord_organolithium_reactions_rich.csv"
USPTO_RICH_CSV = ROOT / "data" / "input" / "uspto_organolithium_extracted_rich.csv"
OUTPUT_DIR = ROOT / "data" / "final_output" / "dataset_comparison"
ENRICHED_XLSX = OUTPUT_DIR / "ALL_organolithium_normalized_20260327_1022_enriched.xlsx"
ENRICHMENT_SUMMARY_JSON = OUTPUT_DIR / "my_dataset_enrichment_summary.json"
COMPARISON_SUMMARY_CSV = OUTPUT_DIR / "organolithium_dataset_comparison_summary.csv"
COMPARISON_NOTES_MD = OUTPUT_DIR / "organolithium_dataset_comparison_notes.md"
COMPLETENESS_PNG = OUTPUT_DIR / "organolithium_dataset_completeness.png"
DISTRIBUTIONS_PNG = OUTPUT_DIR / "organolithium_dataset_distributions.png"


def canonicalize_smiles(smiles: str | None) -> str | None:
    if not smiles or not str(smiles).strip():
        return None
    mol = Chem.MolFromSmiles(str(smiles).strip())
    if mol is None:
        return None
    return Chem.MolToSmiles(mol, canonical=True)


def normalize_name(name: str | None) -> str:
    if not name:
        return ""
    text = str(name).strip().lower()
    text = re.sub(r"[\s\-]+", "", text)
    return text


def build_unique_name_maps_from_user_sheet() -> tuple[dict[str, str], dict[str, str]]:
    wb = openpyxl.load_workbook(USER_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {name: i for i, name in enumerate(header)}

    name_to_smiles = defaultdict(set)
    for row in rows:
        for name_col, smiles_col in [
            ("reactant1_name", "reactant1_smiles"),
            ("reactant2_name", "reactant2_smiles"),
            ("product_name", "product_smiles"),
        ]:
            name = row[idx[name_col]]
            smiles = row[idx[smiles_col]]
            if name not in (None, "") and smiles not in (None, ""):
                can = canonicalize_smiles(smiles)
                if can:
                    name_to_smiles[str(name).strip()].add(can)

    exact = {name: next(iter(values)) for name, values in name_to_smiles.items() if len(values) == 1}
    normalized_bucket = defaultdict(set)
    for name, smiles in exact.items():
        normalized_bucket[normalize_name(name)].add(smiles)
    normalized = {
        key: next(iter(values))
        for key, values in normalized_bucket.items()
        if key and len(values) == 1
    }
    return exact, normalized


def build_unique_name_maps_from_ord() -> tuple[dict[str, str], dict[str, str]]:
    name_to_smiles = defaultdict(set)
    with ORD_RICH_CSV.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            inputs = json.loads(row["inputs_json"]) if row["inputs_json"] else {}
            for reaction_input in inputs.values():
                for component in reaction_input.get("components", []):
                    names = []
                    smiles = []
                    for ident in component.get("identifiers", []):
                        ident_type = ident.get("type")
                        value = ident.get("value")
                        if not value:
                            continue
                        if ident_type in {"NAME", "IUPAC_NAME"}:
                            names.append(value.strip())
                        elif ident_type in {"SMILES", "CXSMILES"}:
                            can = canonicalize_smiles(value)
                            if can:
                                smiles.append(can)
                    for name in names:
                        for smi in smiles:
                            name_to_smiles[name].add(smi)

            outcomes = json.loads(row["outcomes_json"]) if row["outcomes_json"] else []
            for outcome in outcomes:
                for product in outcome.get("products", []):
                    names = []
                    smiles = []
                    for ident in product.get("identifiers", []):
                        ident_type = ident.get("type")
                        value = ident.get("value")
                        if not value:
                            continue
                        if ident_type in {"NAME", "IUPAC_NAME"}:
                            names.append(value.strip())
                        elif ident_type in {"SMILES", "CXSMILES"}:
                            can = canonicalize_smiles(value)
                            if can:
                                smiles.append(can)
                    for name in names:
                        for smi in smiles:
                            name_to_smiles[name].add(smi)

    exact = {name: next(iter(values)) for name, values in name_to_smiles.items() if len(values) == 1}
    normalized_bucket = defaultdict(set)
    for name, smiles in exact.items():
        normalized_bucket[normalize_name(name)].add(smiles)
    normalized = {
        key: next(iter(values))
        for key, values in normalized_bucket.items()
        if key and len(values) == 1
    }
    return exact, normalized


def choose_fill(name: str | None, existing_smiles: str | None, maps: dict[str, dict[str, str]]) -> tuple[str | None, str | None]:
    if existing_smiles not in (None, ""):
        can = canonicalize_smiles(existing_smiles)
        return can or str(existing_smiles), "existing_canonicalized" if can else "existing_original"
    if name in (None, ""):
        return None, None
    key = str(name).strip()
    nkey = normalize_name(key)

    if key in maps["user_exact"]:
        return maps["user_exact"][key], "user_exact_unique_name"
    if nkey in maps["user_normalized"]:
        return maps["user_normalized"][nkey], "user_normalized_unique_name"
    if key in maps["ord_exact"]:
        return maps["ord_exact"][key], "ord_exact_unique_name"
    if nkey in maps["ord_normalized"]:
        return maps["ord_normalized"][nkey], "ord_normalized_unique_name"
    return None, None


def build_reaction_smiles(row: dict) -> tuple[str | None, str]:
    reactants = []
    completeness_parts = []
    for i in (1, 2):
        name = row.get(f"reactant{i}_name")
        smiles = row.get(f"reactant{i}_smiles_enriched")
        if name not in (None, ""):
            completeness_parts.append(f"r{i}")
            if smiles:
                reactants.append(smiles)
            else:
                return None, "missing_named_reactant_smiles"
    product = row.get("product_smiles_enriched")
    if row.get("product_name") not in (None, "") and not product:
        return None, "missing_product_smiles"
    if not reactants or not product:
        return None, "insufficient_structures"
    return ".".join(reactants) + ">>" + product, "+".join(completeness_parts) + "+p"


def enrich_user_dataset() -> pd.DataFrame:
    user_exact, user_normalized = build_unique_name_maps_from_user_sheet()
    ord_exact, ord_normalized = build_unique_name_maps_from_ord()
    maps = {
        "user_exact": user_exact,
        "user_normalized": user_normalized,
        "ord_exact": ord_exact,
        "ord_normalized": ord_normalized,
    }

    df = pd.read_excel(USER_XLSX, sheet_name="All Records")
    for col in ["reactant1_smiles", "reactant2_smiles", "product_smiles", "reaction_smiles"]:
        if col in df.columns:
            df[col] = df[col].where(pd.notna(df[col]), None)

    fill_counter = Counter()
    for entity in ["reactant1", "reactant2", "product"]:
        enriched = []
        sources = []
        for _, row in df.iterrows():
            smi, src = choose_fill(row.get(f"{entity}_name"), row.get(f"{entity}_smiles"), maps)
            enriched.append(smi)
            sources.append(src)
            if src and not src.startswith("existing"):
                fill_counter[f"{entity}:{src}"] += 1
        df[f"{entity}_smiles_enriched"] = enriched
        df[f"{entity}_smiles_enrichment_source"] = sources
        df[f"{entity}_smiles_canonical"] = [canonicalize_smiles(x) if x else None for x in enriched]

    reaction_smiles_enriched = []
    reaction_smiles_status = []
    for _, row in df.iterrows():
        existing = row.get("reaction_smiles")
        if existing not in (None, ""):
            reaction_smiles_enriched.append(existing)
            reaction_smiles_status.append("existing")
            continue
        built, status = build_reaction_smiles(row)
        reaction_smiles_enriched.append(built)
        reaction_smiles_status.append(status)
        if built:
            fill_counter["reaction_smiles:built_from_available_structures"] += 1
    df["reaction_smiles_enriched"] = reaction_smiles_enriched
    df["reaction_smiles_enrichment_status"] = reaction_smiles_status

    summary = {
        "input_file": str(USER_XLSX),
        "output_file": str(ENRICHED_XLSX),
        "rows": int(len(df)),
        "fills_by_source": dict(fill_counter),
        "coverage_before": {
            "reactant1_smiles": int(df["reactant1_smiles"].notna().sum()),
            "reactant2_smiles": int(df["reactant2_smiles"].notna().sum()),
            "product_smiles": int(df["product_smiles"].notna().sum()),
            "reaction_smiles": int(df["reaction_smiles"].notna().sum()),
        },
        "coverage_after": {
            "reactant1_smiles_enriched": int(df["reactant1_smiles_enriched"].notna().sum()),
            "reactant2_smiles_enriched": int(df["reactant2_smiles_enriched"].notna().sum()),
            "product_smiles_enriched": int(df["product_smiles_enriched"].notna().sum()),
            "reaction_smiles_enriched": int(df["reaction_smiles_enriched"].notna().sum()),
        },
    }
    ENRICHMENT_SUMMARY_JSON.write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    df.to_excel(ENRICHED_XLSX, index=False)
    return df


def ord_extract_temperature_value(conditions_json: str | None) -> float | None:
    if not conditions_json:
        return None
    try:
        data = json.loads(conditions_json)
    except Exception:
        return None
    temp = data.get("temperature", {})
    for key in ("setpoint", "control"):
        node = temp.get(key)
        if isinstance(node, dict) and "value" in node and isinstance(node["value"], (int, float)):
            return float(node["value"])
    return None


def ord_extract_time_seconds(outcomes_json: str | None) -> float | None:
    if not outcomes_json:
        return None
    try:
        outcomes = json.loads(outcomes_json)
    except Exception:
        return None
    unit_to_s = {
        "SECOND": 1.0,
        "SECONDS": 1.0,
        "MINUTE": 60.0,
        "MINUTES": 60.0,
        "HOUR": 3600.0,
        "HOURS": 3600.0,
    }
    for outcome in outcomes:
        rt = outcome.get("reaction_time")
        if isinstance(rt, dict) and isinstance(rt.get("value"), (int, float)):
            factor = unit_to_s.get(rt.get("units", ""), None)
            if factor:
                return float(rt["value"]) * factor
    return None


def ord_extract_has_yield(outcomes_json: str | None) -> bool:
    if not outcomes_json:
        return False
    try:
        outcomes = json.loads(outcomes_json)
    except Exception:
        return False
    for outcome in outcomes:
        for product in outcome.get("products", []):
            for measurement in product.get("measurements", []):
                if measurement.get("type") == "YIELD":
                    return True
    return False


def summarize_user_dataset(df: pd.DataFrame) -> dict:
    def present(series_name: str) -> int:
        return int(df[series_name].notna().sum())

    has_outcome = (
        df["yield_pct"].notna()
        | df["batch_yield_pct"].notna()
        | df["conversion_pct"].notna()
        | df["selectivity_pct"].notna()
    )
    structure_complete = (
        df["reaction_smiles_enriched"].notna()
        | (df["reactant1_smiles_enriched"].notna() & df["product_smiles_enriched"].notna())
    )
    return {
        "dataset": "My 130-paper dataset",
        "rows": int(len(df)),
        "unique_sources": int(df["paper_basename"].nunique(dropna=True)),
        "unique_reactions": int(df["reaction_smiles_enriched"].nunique(dropna=True)),
        "structure_complete_pct": float(structure_complete.mean() * 100),
        "reaction_smiles_pct": float(df["reaction_smiles_enriched"].notna().mean() * 100),
        "outcome_pct": float(has_outcome.mean() * 100),
        "yield_pct": float(df["yield_pct"].notna().mean() * 100),
        "temperature_pct": float(df["temperature_C"].notna().mean() * 100),
        "time_pct": float(df["residence_time_s"].notna().mean() * 100),
        "solvent_pct": float(df["solvent"].notna().mean() * 100),
        "catalyst_pct": float(df["catalyst"].notna().mean() * 100),
        "flow_pct": float(
            (
                df["flow_rate_mL_min"].notna()
                | df["flow_rate_stream1_mL_min"].notna()
                | df["flow_rate_stream2_mL_min"].notna()
            ).mean()
            * 100
        ),
        "notes_pct": float(df["notes"].notna().mean() * 100),
    }


def summarize_uspto_dataset() -> tuple[dict, pd.DataFrame]:
    df = pd.read_csv(USPTO_RICH_CSV)
    structure_complete = df["reaction_smiles_raw"].notna() | (
        df["reactant1_smiles"].notna() & df["product1_smiles"].notna()
    )
    summary = {
        "dataset": "USPTO-LLM organolithium",
        "rows": int(len(df)),
        "unique_sources": int(df["source_file"].nunique(dropna=True)),
        "unique_reactions": int(df["reaction_id"].nunique(dropna=True)),
        "structure_complete_pct": float(structure_complete.mean() * 100),
        "reaction_smiles_pct": float(df["reaction_smiles_raw"].notna().mean() * 100),
        "outcome_pct": 0.0,
        "yield_pct": 0.0,
        "temperature_pct": float(df["temperature_C"].notna().mean() * 100),
        "time_pct": float(df["reaction_time_s"].notna().mean() * 100),
        "solvent_pct": float(df["solvent"].notna().mean() * 100),
        "catalyst_pct": float(df["catalyst"].notna().mean() * 100),
        "flow_pct": 0.0,
        "notes_pct": 0.0,
    }
    return summary, df


def summarize_ord_dataset() -> tuple[dict, pd.DataFrame]:
    df = pd.read_csv(ORD_RICH_CSV)
    temp_values = df["conditions_json"].map(ord_extract_temperature_value)
    time_values = df["outcomes_json"].map(ord_extract_time_seconds)
    has_yield = df["outcomes_json"].map(ord_extract_has_yield)
    structure_complete = df["reaction_smiles"].notna()
    summary = {
        "dataset": "ORD organolithium",
        "rows": int(len(df)),
        "unique_sources": int(df["doi"].nunique(dropna=True)),
        "unique_reactions": int(df["reaction_id"].nunique(dropna=True)),
        "structure_complete_pct": float(structure_complete.mean() * 100),
        "reaction_smiles_pct": float(df["reaction_smiles"].notna().mean() * 100),
        "outcome_pct": float(df["has_outcomes"].fillna(False).mean() * 100),
        "yield_pct": float(has_yield.mean() * 100),
        "temperature_pct": float(temp_values.notna().mean() * 100),
        "time_pct": float(time_values.notna().mean() * 100),
        "solvent_pct": float(df["inputs_json"].str.contains('"reaction_role": "SOLVENT"', na=False).mean() * 100),
        "catalyst_pct": float(df["inputs_json"].str.contains('"reaction_role": "CATALYST"', na=False).mean() * 100),
        "flow_pct": float(df["conditions_json"].str.contains('"flow"', na=False).mean() * 100),
        "notes_pct": float(df["has_notes"].fillna(False).mean() * 100),
    }
    df = df.copy()
    df["temperature_C_extracted"] = temp_values
    df["time_s_extracted"] = time_values
    return summary, df


def save_comparison_outputs(user_df: pd.DataFrame, uspto_df: pd.DataFrame, ord_df: pd.DataFrame) -> None:
    user_summary = summarize_user_dataset(user_df)
    uspto_summary, _ = summarize_uspto_dataset()
    ord_summary, _ = summarize_ord_dataset()
    summary_df = pd.DataFrame([user_summary, uspto_summary, ord_summary])
    summary_df.to_csv(COMPARISON_SUMMARY_CSV, index=False)

    notes = [
        "# Organolithium Dataset Comparison",
        "",
        "- My dataset is the only one with strong flow-condition coverage and substantial outcome coverage.",
        "- USPTO-LLM has broader reaction coverage but essentially no outcome labels.",
        "- ORD has rich metadata and notes, but very sparse explicit yield coverage in the matched organolithium subset.",
        "",
        "Files:",
        f"- Enriched workbook: {ENRICHED_XLSX}",
        f"- Enrichment summary: {ENRICHMENT_SUMMARY_JSON}",
        f"- Comparison summary: {COMPARISON_SUMMARY_CSV}",
        f"- Completeness plot: {COMPLETENESS_PNG}",
        f"- Distribution plot: {DISTRIBUTIONS_PNG}",
    ]
    COMPARISON_NOTES_MD.write_text("\n".join(notes), encoding="utf-8")

    plot_df = summary_df.melt(
        id_vars="dataset",
        value_vars=[
            "structure_complete_pct",
            "yield_pct",
            "temperature_pct",
            "time_pct",
            "solvent_pct",
            "catalyst_pct",
            "flow_pct",
            "notes_pct",
        ],
        var_name="metric",
        value_name="percent",
    )
    plt.figure(figsize=(12, 6))
    sns.barplot(data=plot_df, x="metric", y="percent", hue="dataset")
    plt.xticks(rotation=30, ha="right")
    plt.ylabel("Coverage (%)")
    plt.xlabel("")
    plt.title("Information Completeness Across Organolithium Datasets")
    plt.tight_layout()
    plt.savefig(COMPLETENESS_PNG, dpi=200)
    plt.close()

    user_temp = pd.to_numeric(user_df["temperature_C"], errors="coerce").dropna()
    user_time = pd.to_numeric(user_df["residence_time_s"], errors="coerce").dropna()
    user_yield = pd.to_numeric(user_df["yield_pct"], errors="coerce").dropna()
    uspto_temp = pd.to_numeric(uspto_df["temperature_C"], errors="coerce").dropna()
    uspto_time = pd.to_numeric(uspto_df["reaction_time_s"], errors="coerce").dropna()
    ord_temp = pd.to_numeric(ord_df["temperature_C_extracted"], errors="coerce").dropna()
    ord_time = pd.to_numeric(ord_df["time_s_extracted"], errors="coerce").dropna()

    fig, axes = plt.subplots(1, 3, figsize=(15, 4))
    sns.histplot(user_temp, color="#1f77b4", label="Mine", ax=axes[0], bins=40, stat="density", element="step", fill=False)
    sns.histplot(uspto_temp, color="#ff7f0e", label="USPTO-LLM", ax=axes[0], bins=40, stat="density", element="step", fill=False)
    sns.histplot(ord_temp, color="#2ca02c", label="ORD", ax=axes[0], bins=40, stat="density", element="step", fill=False)
    axes[0].set_title("Temperature Distribution")
    axes[0].set_xlabel("Temperature (C)")
    axes[0].legend()

    sns.histplot(user_time[user_time > 0], color="#1f77b4", label="Mine", ax=axes[1], bins=40, stat="density", element="step", fill=False)
    sns.histplot(uspto_time[uspto_time > 0], color="#ff7f0e", label="USPTO-LLM", ax=axes[1], bins=40, stat="density", element="step", fill=False)
    sns.histplot(ord_time[ord_time > 0], color="#2ca02c", label="ORD", ax=axes[1], bins=40, stat="density", element="step", fill=False)
    axes[1].set_xscale("log")
    axes[1].set_title("Time Distribution")
    axes[1].set_xlabel("Time (s, log scale)")
    axes[1].legend()

    sns.histplot(user_yield, color="#1f77b4", ax=axes[2], bins=40)
    axes[2].set_title("My Dataset Yield Distribution")
    axes[2].set_xlabel("yield_pct")
    plt.tight_layout()
    plt.savefig(DISTRIBUTIONS_PNG, dpi=200)
    plt.close()


def main() -> None:
    RDLogger.DisableLog("rdApp.*")
    sns.set_theme(style="whitegrid")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    user_df = enrich_user_dataset()
    uspto_df = pd.read_csv(USPTO_RICH_CSV)
    ord_df = pd.read_csv(ORD_RICH_CSV)
    ord_df["temperature_C_extracted"] = ord_df["conditions_json"].map(ord_extract_temperature_value)
    ord_df["time_s_extracted"] = ord_df["outcomes_json"].map(ord_extract_time_seconds)

    save_comparison_outputs(user_df, uspto_df, ord_df)

    print(f"Saved enriched workbook: {ENRICHED_XLSX}")
    print(f"Saved enrichment summary: {ENRICHMENT_SUMMARY_JSON}")
    print(f"Saved comparison summary: {COMPARISON_SUMMARY_CSV}")
    print(f"Saved notes: {COMPARISON_NOTES_MD}")
    print(f"Saved completeness plot: {COMPLETENESS_PNG}")
    print(f"Saved distributions plot: {DISTRIBUTIONS_PNG}")


if __name__ == "__main__":
    main()
